from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(r"C:\Users\dungl\Downloads\NCKH")
SOURCE = ROOT / "_bmad-output" / "planning-artifacts" / "danh-muc-use-case-chi-tiet-va-doi-chieu-tong-quan.md"
OUTPUT = ROOT / "_bmad-output" / "planning-artifacts" / "danh-muc-use-case-tong-quan-chi-tiet-va-ma-tran-actor.docx"
SKILL_SCRIPTS = Path(r"C:\Users\dungl\.codex\plugins\cache\openai-primary-runtime\documents\26.715.12143\skills\documents\scripts")
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa


BLUE = "1F4E78"
BLUE_DARK = "17365D"
BLUE_LIGHT = "E8EEF5"
ROW_ALT = "F7F9FC"
GRID = "B7C9D6"
TEXT = "1F2933"
MUTED = "59626C"

STEP_RULES = [
    ("Bước 01 – Đăng ký đề tài", lambda value: re.match(r"^UC-(DOT|DK)-", value)),
    ("Bước 02 – Phê duyệt sơ bộ", lambda value: re.match(r"^UC-XDSB-", value)),
    ("Bước 03 – Viết và nộp thuyết minh", lambda value: re.match(r"^UC-TM-0[1-4]$", value)),
    ("Bước 04 – Phê duyệt thuyết minh và giao thực hiện đề tài", lambda value: re.match(r"^UC-TM-(0[5-9]|1[0-5])$", value)),
    ("Bước 05 – Thực hiện và báo cáo tiến độ giữa kỳ", lambda value: re.match(r"^UC-TD-", value)),
    ("Bước 06 – Nộp hồ sơ và nghiệm thu đề tài", lambda value: re.match(r"^UC-NT-", value)),
    ("Bước 07 – Chỉnh sửa theo góp ý của Hội đồng nghiệm thu", lambda value: re.match(r"^UC-KQ-0[1-4]$", value)),
    ("Bước 08 – Công nhận kết quả nghiên cứu đề tài", lambda value: re.match(r"^UC-KQ-0[5-6]$", value)),
    ("Bước 09 – Theo dõi triển khai ứng dụng", lambda value: re.match(r"^UC-KQ-07$", value)),
    ("Ngoài quy trình 9 bước – Tra cứu số tiết NCKH", lambda value: re.match(r"^UC-ST-01$", value)),
]


def parse_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse_source():
    text = SOURCE.read_text(encoding="utf-8")
    sec1 = text[text.index("## 1."):text.index("## 2.")]
    sec2 = text[text.index("## 2."):text.index("## 3.")]
    sec3 = text[text.index("## 3."):text.index("## 4.")]
    sec4 = text[text.index("## 4."):]
    hucs = [parse_row(line) for line in sec1.splitlines() if re.match(r"^\| HUC-\d{2} \|", line)]
    details = [parse_row(line) for line in sec2.splitlines() if re.match(r"^\| UC-(?!DL-)[A-Z]+-\d+ \|", line)]
    reusable = [parse_row(line) for line in sec2.splitlines() if re.match(r"^\| UC-DL-\d+ \|", line)]
    matrix = [parse_row(line) for line in sec3.splitlines() if re.match(r"^\| UC-[A-Z]+-\d+ \|", line)]
    notes = [line[2:].strip() for line in sec4.splitlines() if line.startswith("- ")]
    assert len(hucs) == 16
    assert len(details) == 72
    assert len(reusable) == 5
    assert len(matrix) == 72
    assert len(notes) == 5
    assert {row[0] for row in details} == {row[0] for row in matrix}
    return hucs, details, reusable, matrix, notes


def set_run_font(run, size, color=TEXT, bold=False, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rpr.rFonts.set(qn(key), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color=GRID, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    item = tc_pr.find(qn("w:tcBorders"))
    if item is None:
        item = OxmlElement("w:tcBorders")
        tc_pr.append(item)
    for edge in ("top", "left", "bottom", "right"):
        node = item.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            item.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def fill_cell(cell, text, size=8, header=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    if text == "✓":
        set_run_font(run, 10, BLUE_DARK, True, "Segoe UI Symbol")
    else:
        set_run_font(run, size, "FFFFFF" if header else TEXT, header)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, 8, MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    return p


def add_table(doc, headers, rows, weights, sizes, centers):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    repeat_header(table.rows[0])
    prevent_split(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, BLUE)
        borders(cell, BLUE_DARK, "5")
        fill_cell(cell, header, 7.3, True, True)
    for row_index, values in enumerate(rows, start=1):
        row = table.add_row()
        prevent_split(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            borders(cell)
            if row_index % 2 == 0:
                shade(cell, ROW_ALT)
            fill_cell(cell, value, sizes[col_index], False, centers[col_index])
            if col_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    section = doc.sections[-1]
    total = section_content_width_dxa(section) - 70
    widths = column_widths_from_weights(weights, total)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=total,
        indent_dxa=70,
        cell_margins_dxa={"top": 55, "bottom": 55, "start": 70, "end": 70},
    )
    return table


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(9)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h1.font.size = Pt(15)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(BLUE)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h2.font.size = Pt(11.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    h2.paragraph_format.space_before = Pt(9)
    h2.paragraph_format.space_after = Pt(4)


def build():
    hucs, details, reusable, matrix, notes = parse_source()
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Mm(10))
    section.header_distance = Mm(6)
    section.footer_distance = Mm(6)
    configure_styles(doc)

    hp = section.header.paragraphs[0]
    hp.text = "Danh mục use case | Hệ thống quản lý hoạt động NCKH cấp trường"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in hp.runs:
        set_run_font(run, 8, MUTED)
    add_page_field(section.footer.paragraphs[0])

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run("DANH MỤC USE CASE HỆ THỐNG NCKH"), 19, BLUE, True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    subtitle.paragraph_format.keep_with_next = True
    set_run_font(subtitle.add_run("Use case tổng quan, use case chi tiết theo quy trình và ma trận actor"), 10, MUTED)

    add_heading(doc, "1. Danh mục use case tổng quan", 1)
    add_table(
        doc,
        ["Mã", "Use case tổng quan", "Actor"],
        hucs,
        [0.8, 5.4, 4.4],
        [8.2, 8.5, 8.3],
        [True, False, False],
    )

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "2. Danh mục use case chi tiết theo quy trình", 1)
    assigned = set()
    for step_name, predicate in STEP_RULES:
        step_rows = [row for row in details if predicate(row[0])]
        assert step_rows, step_name
        assigned.update(row[0] for row in step_rows)
        add_heading(doc, step_name, 2)
        add_table(
            doc,
            ["Mã", "Use case chi tiết", "UC tổng quan cha", "Tác nhân chính", "Kết quả/đầu ra"],
            step_rows,
            [0.85, 2.7, 1.35, 2.35, 3.25],
            [7.3, 7.8, 7.3, 7.5, 7.5],
            [True, False, True, False, False],
        )
    assert len(assigned) == 72

    add_heading(doc, "Năng lực dùng chung", 2)
    add_table(
        doc,
        ["Mã", "Năng lực dùng chung", "Được sử dụng trong", "Quy tắc"],
        reusable,
        [0.9, 3.0, 3.4, 3.2],
        [7.5, 8.0, 7.7, 7.7],
        [True, False, False, False],
    )

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "3. Ma trận use case – actor", 1)
    matrix_headers = [
        "Mã", "Use case chi tiết", "Giảng viên", "Sinh viên",
        "Cán bộ phụ trách Khoa", "Cán bộ phụ trách P.KHCN",
        "Thành viên HĐ xét duyệt hồ sơ", "Thành viên HĐ xét duyệt thuyết minh",
        "Thành viên HĐ nghiệm thu đề tài",
    ]
    add_table(
        doc,
        matrix_headers,
        matrix,
        [0.72, 3.58, 0.75, 0.75, 0.9, 0.9, 0.95, 1.0, 1.0],
        [7.2, 7.8, 8.0, 8.0, 7.2, 7.2, 7.0, 7.0, 7.0],
        [True, False, True, True, True, True, True, True, True],
    )

    add_heading(doc, "4. Ghi chú phạm vi", 1)
    for note in notes:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        set_run_font(p.add_run(note), 9, TEXT)

    doc.core_properties.title = "Danh mục use case hệ thống NCKH"
    doc.core_properties.subject = "Use case tổng quan, chi tiết và ma trận actor"
    doc.core_properties.author = "Mary – Business Analyst"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"HUC={len(hucs)} Detail={len(details)} Matrix={len(matrix)}")


if __name__ == "__main__":
    build()
