from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


PROJECT = Path(r"C:\Users\dungl\Downloads\NCKH")
SOURCE = PROJECT / "_bmad-output" / "planning-artifacts" / "danh-muc-use-case-chi-tiet-va-doi-chieu-tong-quan.md"
OUTPUT = PROJECT / "_bmad-output" / "planning-artifacts" / "ma-tran-use-case-actor.docx"
SKILL_SCRIPTS = Path(r"C:\Users\dungl\.codex\plugins\cache\openai-primary-runtime\documents\26.715.12143\skills\documents\scripts")
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa


ACTORS = [
    "Giảng viên",
    "Sinh viên",
    "Cán bộ phụ trách Khoa",
    "Cán bộ phụ trách P.KHCN",
    "Thành viên HĐ xét duyệt hồ sơ",
    "Thành viên HĐ xét duyệt thuyết minh",
    "Thành viên HĐ nghiệm thu đề tài",
]

HEADER_LABELS = [
    "Mã",
    "Use case chi tiết",
    "Giảng viên",
    "Sinh viên",
    "Cán bộ phụ trách\nKhoa",
    "Cán bộ phụ trách\nP.KHCN",
    "Thành viên HĐ xét duyệt\nhồ sơ",
    "Thành viên HĐ xét duyệt\nthuyết minh",
    "Thành viên HĐ nghiệm thu\nđề tài",
]

BLUE = "1F4E78"
BLUE_DARK = "17365D"
BLUE_LIGHT = "E8EEF5"
ROW_ALT = "F7F9FC"
GRID = "B7C9D6"
MUTED = RGBColor(89, 98, 108)


def parse_matrix():
    text = SOURCE.read_text(encoding="utf-8")
    start = text.index("## 3. Ma trận use case – actor")
    end = text.index("## 4. Ghi chú phạm vi")
    rows = []
    for line in text[start:end].splitlines():
        if not re.match(r"^\| UC-[A-Z]+-\d+ \|", line):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) != 9:
            raise ValueError(f"Dòng ma trận không có 9 cột: {line}")
        rows.append(cells)
    if len(rows) != 72:
        raise ValueError(f"Cần đúng 72 use case, tìm thấy {len(rows)}")
    ids = [row[0] for row in rows]
    if len(set(ids)) != 72:
        raise ValueError("Mã use case trong ma trận bị trùng")
    for row in rows:
        if not any(cell == "✓" for cell in row[2:]):
            raise ValueError(f"Use case chưa có actor: {row[0]}")
    return rows


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=GRID, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:color"), color)


def set_run_font(run, name="Calibri", size=8, color=None, bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def format_cell_text(cell, text, *, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    parts = text.split("\n")
    for idx, part in enumerate(parts):
        if idx:
            p.add_run().add_break()
        run = p.add_run(part)
        if text == "✓":
            set_run_font(run, "Segoe UI Symbol", 10, BLUE_DARK, True)
        elif header:
            set_run_font(run, "Calibri", 7.2, "FFFFFF", True)
        else:
            set_run_font(run, "Calibri", 8.0, "1F2933", False)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_repeat(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, "Calibri", 8, "6B7280", False)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def build():
    rows = parse_matrix()
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(10)
    section.bottom_margin = Mm(10)
    section.left_margin = Mm(10)
    section.right_margin = Mm(10)
    section.header_distance = Mm(6)
    section.footer_distance = Mm(6)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    header_p = section.header.paragraphs[0]
    header_p.text = "Ma trận use case – actor | Hệ thống quản lý hoạt động NCKH cấp trường"
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    for run in header_p.runs:
        set_run_font(run, "Calibri", 8, "6B7280", False)

    footer_p = section.footer.paragraphs[0]
    add_page_field(footer_p)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.keep_with_next = True
    title_run = title.add_run("MA TRẬN USE CASE – ACTOR")
    set_run_font(title_run, "Calibri", 18, BLUE, True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(6)
    subtitle.paragraph_format.keep_with_next = True
    subtitle_run = subtitle.add_run("Hệ thống quản lý hoạt động nghiên cứu khoa học cấp trường · 72 use case · 7 actor")
    set_run_font(subtitle_run, "Calibri", 9, None, False)
    subtitle_run.font.color.rgb = MUTED

    table = doc.add_table(rows=1, cols=9)
    table.style = "Table Grid"
    table.autofit = False

    header = table.rows[0]
    mark_header_repeat(header)
    prevent_row_split(header)
    for idx, label in enumerate(HEADER_LABELS):
        cell = header.cells[idx]
        set_cell_shading(cell, BLUE)
        set_cell_borders(cell, BLUE_DARK, "5")
        format_cell_text(cell, label, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row_index, source_row in enumerate(rows, start=1):
        row = table.add_row()
        prevent_row_split(row)
        for col_index, value in enumerate(source_row):
            cell = row.cells[col_index]
            set_cell_borders(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, ROW_ALT)
            align = WD_ALIGN_PARAGRAPH.LEFT if col_index == 1 else WD_ALIGN_PARAGRAPH.CENTER
            format_cell_text(cell, value, align=align)
            if col_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(BLUE_DARK)

    content_width = section_content_width_dxa(section)
    table_width = content_width - 70
    weights = [0.72, 3.58, 0.75, 0.75, 0.90, 0.90, 0.95, 1.00, 1.00]
    widths = column_widths_from_weights(weights, table_width)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=table_width,
        indent_dxa=70,
        cell_margins_dxa={"top": 55, "bottom": 55, "start": 70, "end": 70},
    )

    doc.core_properties.title = "Ma trận use case – actor"
    doc.core_properties.subject = "Hệ thống quản lý hoạt động nghiên cứu khoa học cấp trường"
    doc.core_properties.author = "Mary – Business Analyst"
    doc.core_properties.keywords = "use case, actor, NCKH"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Rows: {len(rows)}")


if __name__ == "__main__":
    build()
