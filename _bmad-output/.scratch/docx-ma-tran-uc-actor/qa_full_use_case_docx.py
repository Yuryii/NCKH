from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\dungl\Downloads\NCKH")
SOURCE = ROOT / "_bmad-output" / "planning-artifacts" / "danh-muc-use-case-chi-tiet-va-doi-chieu-tong-quan.md"
DOCX = ROOT / "_bmad-output" / "planning-artifacts" / "danh-muc-use-case-tong-quan-chi-tiet-va-ma-tran-actor.docx"


def rows_from(block, pattern):
    result = []
    for line in block.splitlines():
        if re.match(pattern, line):
            result.append([cell.strip() for cell in line.strip().strip("|").split("|")])
    return result


text = SOURCE.read_text(encoding="utf-8")
sec1 = text[text.index("## 1."):text.index("## 2.")]
sec2 = text[text.index("## 2."):text.index("## 3.")]
sec3 = text[text.index("## 3."):text.index("## 4.")]
hucs = rows_from(sec1, r"^\| HUC-\d{2} \|")
details = rows_from(sec2, r"^\| UC-(?!DL-)[A-Z]+-\d+ \|")
reusable = rows_from(sec2, r"^\| UC-DL-\d+ \|")
matrix = rows_from(sec3, r"^\| UC-[A-Z]+-\d+ \|")

doc = Document(DOCX)
assert len(doc.tables) == 13, f"Expected 13 tables, got {len(doc.tables)}"
assert [len(t.rows) for t in doc.tables] == [17, 22, 11, 5, 12, 7, 13, 5, 3, 2, 2, 6, 73]
assert [len(t.columns) for t in doc.tables] == [3] + [5] * 10 + [4, 9]

doc_hucs = [[cell.text.strip() for cell in row.cells] for row in doc.tables[0].rows[1:]]
assert doc_hucs == hucs

doc_details = []
for table in doc.tables[1:11]:
    doc_details.extend([[cell.text.strip() for cell in row.cells] for row in table.rows[1:]])
assert doc_details == details

doc_reusable = [[cell.text.strip() for cell in row.cells] for row in doc.tables[11].rows[1:]]
assert doc_reusable == reusable

doc_matrix = [[cell.text.strip() for cell in row.cells] for row in doc.tables[12].rows[1:]]
assert doc_matrix == matrix
assert all(any(value == "✓" for value in row[2:]) for row in doc_matrix)

headings = [p.text.strip() for p in doc.paragraphs if p.style.name.startswith("Heading")]
for step in (
    "Bước 01 – Đăng ký đề tài", "Bước 02 – Phê duyệt sơ bộ",
    "Bước 03 – Viết và nộp thuyết minh",
    "Bước 04 – Phê duyệt thuyết minh và giao thực hiện đề tài",
    "Bước 05 – Thực hiện và báo cáo tiến độ giữa kỳ",
    "Bước 06 – Nộp hồ sơ và nghiệm thu đề tài",
    "Bước 07 – Chỉnh sửa theo góp ý của Hội đồng nghiệm thu",
    "Bước 08 – Công nhận kết quả nghiên cứu đề tài",
    "Bước 09 – Theo dõi triển khai ứng dụng",
    "Ngoài quy trình 9 bước – Tra cứu số tiết NCKH",
):
    assert step in headings, f"Missing heading: {step}"

section = doc.sections[0]
assert section.orientation == WD_ORIENT.LANDSCAPE and section.page_width > section.page_height
for table in doc.tables:
    assert table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
    assert all(row.height is None for row in table.rows)

print("PASS: 16 HUC, 72 detailed UC, 5 reusable capabilities, 72 matrix rows")
print("PASS: 9 process steps plus the out-of-process UC-ST-01 group")
print("PASS: 13 tables match Markdown source exactly")
print("PASS: A4 landscape, repeating headers, no fixed row heights")
