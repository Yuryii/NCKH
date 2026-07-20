import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


DOC = Path(r"C:\Users\dungl\Downloads\NCKH\_bmad-output\FRS Hệ thống quản lý hoạt động NCKH cấp trường - Dự thảo.docx")
d = Document(DOC)

errors = []
all_text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)

for banned in ["ngân hàng câu hỏi", "ghép đề thi", "tổ chức thi", "chấm điểm tự động"]:
    if banned.lower() in all_text.lower():
        errors.append(f"Banned residual text: {banned}")

req_codes = []
for t in d.tables:
    headers = [c.text.strip() for c in t.rows[0].cells]
    if headers[:2] == ["Mã YC", "Tên chức năng"]:
        req_codes.extend(r.cells[0].text.strip() for r in t.rows[1:] if re.fullmatch(r"UC-[A-Z]+-\d+", r.cells[0].text.strip()))

if len(req_codes) != 73:
    errors.append(f"Requirement count {len(req_codes)} != 73")
if len(set(req_codes)) != 73:
    errors.append("Duplicate requirement codes")
if "UC-TR-01" in req_codes:
    errors.append("UC-TR-01 incorrectly included in implementation requirements")
if "UC-TR-01" not in all_text:
    errors.append("UC-TR-01 missing from open issues")

us_codes = []
for t in d.tables:
    headers = [c.text.strip() for c in t.rows[0].cells]
    if headers[:2] == ["Mã US", "UCTQ"]:
        us_codes.extend(r.cells[1].text.strip() for r in t.rows[1:])
if len(us_codes) != 19 or len(set(us_codes)) != 19:
    errors.append(f"UCTQ/US count invalid: {len(us_codes)}")

font_issues = []
normal_size = d.styles["Normal"].font.size.pt if d.styles["Normal"].font.size else None

def effective_size(run, paragraph):
    if run.font.size is not None:
        return run.font.size.pt
    if paragraph.style is not None and paragraph.style.font.size is not None:
        return paragraph.style.font.size.pt
    return normal_size

for p_idx, p in enumerate(d.paragraphs):
    if not p.text.strip() or p.style.name.startswith("Heading"):
        continue
    # Cover roles are intentionally larger than 13 pt.
    sizes = [r.font.size.pt for r in p.runs if r.text.strip() and r.font.size]
    if p_idx <= 20 and any(s > 13.1 for s in sizes):
        continue
    for r in p.runs:
        size = effective_size(r, p)
        if r.text.strip() and (size is None or abs(size - 13) > 0.1):
            font_issues.append(("paragraph", p_idx, r.text[:40], size))

for t_idx, t in enumerate(d.tables):
    for r_idx, row in enumerate(t.rows):
        for c_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for run in p.runs:
                    size = effective_size(run, p)
                    if run.text.strip() and (size is None or abs(size - 13) > 0.1):
                        font_issues.append(("table", (t_idx, r_idx, c_idx), run.text[:40], size))
if font_issues:
    errors.append(f"Font-size issues: {font_issues[:10]} total={len(font_issues)}")

header_flags = 0
nosplit_rows = 0
total_rows = 0
for t in d.tables:
    if t.rows:
        trpr = t.rows[0]._tr.trPr
        if trpr is not None and trpr.find(qn("w:tblHeader")) is not None:
            header_flags += 1
    for row in t.rows:
        total_rows += 1
        trpr = row._tr.trPr
        if trpr is not None and trpr.find(qn("w:cantSplit")) is not None:
            nosplit_rows += 1
if header_flags != len(d.tables):
    errors.append(f"Repeating headers {header_flags}/{len(d.tables)}")
if nosplit_rows != total_rows:
    errors.append(f"No-split rows {nosplit_rows}/{total_rows}")

print({
    "paragraphs": len(d.paragraphs),
    "tables": len(d.tables),
    "sections": len(d.sections),
    "requirements": len(req_codes),
    "uctq_user_stories": len(us_codes),
    "table_rows": total_rows,
    "header_flags": header_flags,
    "no_split_rows": nosplit_rows,
    "errors": errors,
})
sys.exit(1 if errors else 0)
