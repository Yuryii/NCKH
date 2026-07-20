from __future__ import annotations

import re
from collections import defaultdict
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\dungl\Downloads\NCKH")
TEMPLATE = Path(r"C:\Users\dungl\Downloads\FRS_Mau.docx")
UC_DOCX = ROOT / "_bmad-output" / "Danh mục use case chi tiết hệ thống NCKH - có giải thích.docx"
BASELINE = ROOT / "_bmad-output" / "brainstorming" / "brainstorm-liet-ke-day-du-use-case-chi-tiet-nckh-2026-07-18" / "danh-muc-use-case-chi-tiet.md"
OUTPUT = ROOT / "_bmad-output" / "FRS Hệ thống quản lý hoạt động NCKH cấp trường - Dự thảo.docx"

BLUE = "1F4E78"
MID_BLUE = "D9EAF7"
LIGHT_BLUE = "EEF5FA"
GRAY = "E7E6E6"
TEXT = RGBColor(28, 28, 28)
FONT = "Arial"


MODULES = [
    ("M01", "Quản lý đợt đăng ký", 1),
    ("M02", "Đăng ký và xét duyệt tuyến đầu", 2),
    ("M03", "Xét duyệt đề xuất sơ bộ", 3),
    ("M04", "Nộp thuyết minh", 4),
    ("M05", "Phê duyệt thuyết minh và hợp đồng", 5),
    ("M06", "Thực hiện và báo cáo tiến độ", 6),
    ("M07", "Nộp hồ sơ và nghiệm thu", 7),
    ("M08", "Chỉnh sửa sau nghiệm thu và BM14", 8),
    ("M09", "Công nhận kết quả", 9),
    ("M10", "Actor, phân quyền và truy cập", 10),
    ("M11", "Quản lý Hội đồng dùng chung", 11),
    ("M12", "Xử lý biên bản Hội đồng", 12),
    ("M13", "Pipeline biểu mẫu/PDF dùng chung", 13),
    ("M14", "Thông báo và truy vết", 14),
]


MODULE_FIELDS = {
    "M01": ["Mã và tên đợt", "Loại đợt", "Thời gian bắt đầu/kết thúc", "Trạng thái đợt", "Danh mục đề tài giao trực tiếp", "Thời điểm công bố"],
    "M02": ["Loại đề tài và hình thức đăng ký", "Thông tin Chủ nhiệm/người đăng ký", "Giảng viên hướng dẫn hoặc đơn vị", "Nhóm nghiên cứu", "BM01 và phiên bản", "Trạng thái hồ sơ", "Lý do trả/hủy", "Thời hạn cấm đăng ký"],
    "M03": ["Hồ sơ/BM01 được xét", "Hội đồng xét duyệt hồ sơ", "Phiếu BM02", "Biên bản BM03", "Số phiếu hợp lệ", "Kết luận xét duyệt"],
    "M04": ["Mã đề tài", "Loại BM04A/B", "Tệp thuyết minh", "Phiên bản tệp", "Người tải", "Thời điểm tải"],
    "M05": ["Quyết định BM05", "Hội đồng thuyết minh", "Hạn đánh giá", "Phiếu BM06", "Biên bản BM07", "Kết luận thực hiện/không thực hiện", "Hợp đồng đã ký"],
    "M06": ["Đề tài đang thực hiện", "Kỳ báo cáo", "Dữ liệu và PDF BM08", "Tình trạng nộp", "Tài liệu được phép xem/tải"],
    "M07": ["Báo cáo tổng kết BM09", "Quyết định BM10", "Hội đồng nghiệm thu", "Phiếu BM11", "Biên bản BM12", "Kết quả nghiệm thu"],
    "M08": ["Yêu cầu chỉnh sửa sau nghiệm thu", "Giải trình BM13", "Tệp chỉnh sửa", "BM14 hoàn chỉnh", "Thời điểm lưu/công bố"],
    "M09": ["Quyết định BM15", "Danh sách đề tài được công nhận", "Tệp quyết định", "Phạm vi công bố", "Thời điểm công bố"],
    "M10": ["Tài khoản", "Vai trò nghiệp vụ", "Đơn vị", "Hội đồng được phân công", "Thời gian hiệu lực quyền", "Phạm vi dữ liệu"],
    "M11": ["Loại Hội đồng", "Đợt/giai đoạn", "Danh sách thành viên", "Vai trò Chủ tịch/Thư ký/thành viên", "Hạn nộp phiếu", "Trạng thái Hội đồng"],
    "M12": ["Loại biên bản", "Phiên bản biên bản", "Chữ ký Thư ký", "Chữ ký Chủ tịch", "Lý do trả sửa", "Hạn xử lý/gia hạn", "Trạng thái biên bản"],
    "M13": ["Mã biểu mẫu", "Dữ liệu form", "Bản xem trước", "PDF xuất", "PDF/ảnh chữ ký", "Phiên bản", "Người nộp", "Trạng thái khóa"],
    "M14": ["Loại sự kiện", "Người nhận", "Nội dung thông báo", "Thời điểm gửi/xem", "Đối tượng nghiệp vụ", "Audit log trạng thái trước/sau", "Lý do thao tác"],
}


MODULE_RULES = {
    "M01": [
        "Đợt có bốn trạng thái: Nháp, Đã công bố - chưa mở, Đang mở và Đã đóng; không có trạng thái Hủy hoặc Mở lại trong phiên bản đầu.",
        "Hệ thống tự chuyển Đang mở/Đã đóng theo thời gian; P.KHCN không đóng thủ công trước hạn.",
        "Sau khi đợt đã công bố và đã có người tạo hoặc nộp hồ sơ, không được sửa thời gian hoặc danh mục đề tài giao trực tiếp.",
    ],
    "M02": [
        "Hồ sơ chỉ được nộp hợp lệ sau khi có BM01 đã ký; sau khi nộp người đăng ký không tự sửa hoặc rút hồ sơ.",
        "Hồ sơ sinh viên do Giảng viên hướng dẫn duyệt; hồ sơ giảng viên do Trưởng Khoa/Trưởng đơn vị duyệt.",
        "Khi trả sửa phải nhập lý do; nếu nội dung thay đổi, BM01 cũ mất hiệu lực và phải ký/nộp lại trước hạn.",
        "Mỗi người được tham gia tối đa hai đề tài trong cùng đợt; Giảng viên hướng dẫn không bị tính vào hạn mức.",
        "Đề tài giao trực tiếp thuộc về hồ sơ được Trưởng Khoa duyệt sớm nhất; hồ sơ còn lại chuyển Không được chọn.",
    ],
    "M03": [
        "Mọi thành viên đánh giá, gồm Thư ký và P.KHCN khi giữ vai trò Chủ tịch, nộp phiếu BM02 của mình.",
        "Phiếu đã nộp bị khóa; biên bản chỉ được mở sau khi Thư ký chốt tập phiếu theo điều kiện hiện hành.",
        "BM03 do Thư ký lập và P.KHCN/Chủ tịch rà soát, trả sửa hoặc xác nhận theo luồng biên bản dùng chung.",
    ],
    "M04": [
        "BM04 được soạn bên ngoài; hệ thống chỉ lưu tệp hoàn chỉnh theo phiên bản và cung cấp đúng quyền.",
        "Hệ thống không trích xuất hoặc cấu trúc hóa kế hoạch, sản phẩm và nội dung nghiên cứu từ BM04.",
    ],
    "M05": [
        "BM05 được lập/ký bên ngoài và chỉ được P.KHCN đăng, lưu, công bố theo quyền.",
        "BM06 đã nộp là chốt tuyệt đối, không trả lại, sửa hoặc nộp lại; gia hạn chỉ áp dụng người chưa nộp.",
        "Sau khi P.KHCN xác nhận BM07 hợp lệ, kết luận được thông báo; đề tài đạt chuyển sang Đang thực hiện.",
        "Hợp đồng được ký ngoài hệ thống; hệ thống chỉ lưu và cung cấp tệp hợp đồng đã ký.",
    ],
    "M06": [
        "Chủ nhiệm lập/nộp BM08 và theo dõi tài liệu của đúng đề tài theo quyền.",
        "Nguồn kế hoạch/sản phẩm không tự động lấy từ BM04 vì BM04 chỉ là tệp; trường dữ liệu chi tiết cần xác minh theo biểu mẫu.",
    ],
    "M07": [
        "BM09 và BM10 được tải dưới dạng tài liệu hoàn chỉnh; hệ thống không soạn quyết định BM10.",
        "BM11 đã nộp bị khóa và không được sửa/nộp lại.",
        "BM12 tuân theo luồng biên bản; kết quả nghiệm thu chỉ có hiệu lực trên hệ thống sau khi được xác nhận.",
    ],
    "M08": [
        "BM13 do Chủ nhiệm lập/nộp theo yêu cầu chỉnh sửa sau nghiệm thu.",
        "BM14 được lập, xử lý và ký hoàn toàn bên ngoài; P.KHCN chỉ lưu bản hoàn chỉnh và cấp quyền xem.",
    ],
    "M09": [
        "BM15 được lập/ký ngoài hệ thống; P.KHCN chỉ đăng và lưu tệp đã có.",
        "Việc triển khai ứng dụng và lưu hồ sơ chuyên môn phía sau nằm ngoài phạm vi xử lý chi tiết của hệ thống.",
    ],
    "M10": [
        "Người dùng trong Trường dùng tài khoản do Nhà trường quản lý; người ngoài Trường dùng tài khoản do P.KHCN cấp.",
        "Một tài khoản có thể mang nhiều vai trò Hội đồng; quyền được xác định theo vai trò và Hội đồng/giai đoạn.",
        "Giảng viên hướng dẫn là vai trò trên tài khoản Giảng viên, không phải actor/tài khoản riêng.",
    ],
    "M11": [
        "Mỗi giai đoạn có Hội đồng riêng; danh sách thành viên không được sửa sau khi Hội đồng đã được tạo.",
        "Hội đồng có năm trạng thái: Đã thành lập, Đang đánh giá, Đang họp/lập biên bản, Đã hoàn tất và Đã giải tán.",
        "Sau khi giải tán, thành viên chỉ xem lịch sử/chi tiết và không được tải BM04/BM06/BM07 hoặc tài liệu tương ứng.",
    ],
    "M12": [
        "BM03/BM07/BM12 do Thư ký lập; P.KHCN kiểm tra tính hợp lệ và chữ ký, không đánh giá thay Hội đồng.",
        "Khi trả biên bản phải nhập lý do; Thư ký sửa và nộp lại, bản/PDF cũ mất hiệu lực theo quy tắc phiên bản.",
        "Chỉ sau khi P.KHCN xác nhận, hệ thống mới ghi nhận kết quả và gửi thông báo cho actor liên quan.",
    ],
    "M13": [
        "Trước khi bấm Nộp, form còn sửa được; thao tác Nộp làm khóa bản chính thức.",
        "BM02/BM06/BM11 không được sửa hoặc nộp lại sau khi nộp.",
        "BM03/BM07/BM12 có vòng trả chỉnh sửa và trạng thái Chờ P.KHCN xác nhận/Đã xác nhận.",
        "Nếu dữ liệu form và PDF đã ký khác nhau, nguồn chính thức và cách đối soát phải được xác minh trước triển khai.",
    ],
    "M14": [
        "Thông báo được gửi trong ứng dụng theo sự kiện và chỉ đến actor liên quan; chưa dùng email hoặc SMS.",
        "Audit log được lưu theo cơ chế chỉ thêm mới, không cho chức năng nghiệp vụ sửa hoặc xóa bản ghi cũ.",
        "Audit log tối thiểu có actor, thời điểm, hành động, đối tượng, trạng thái trước/sau và lý do nếu có.",
    ],
}


ACTORS = [
    ("Giảng viên", "Đăng ký đề tài giảng viên; thực hiện vai trò Giảng viên hướng dẫn đối với hồ sơ sinh viên được gán."),
    ("Sinh viên", "Tạo, nộp và theo dõi hồ sơ đề tài sinh viên; thực hiện nghĩa vụ của Chủ nhiệm đề tài."),
    ("Trưởng Khoa/Trưởng đơn vị", "Xét duyệt tuyến đầu hồ sơ giảng viên thuộc đơn vị."),
    ("Cán bộ/Phòng KHCN (P.KHCN)", "Quản trị đợt, Hội đồng, tài liệu, kết quả và quyền ngoài Trường; có thể giữ vai trò Chủ tịch Hội đồng."),
    ("Thành viên HĐ xét duyệt hồ sơ", "Xem hồ sơ và nộp phiếu BM02."),
    ("Thư ký HĐ xét duyệt hồ sơ", "Nộp BM02, chốt phiếu và lập BM03."),
    ("Thành viên HĐ xét duyệt thuyết minh", "Xem BM04/BM05 và nộp BM06."),
    ("Thư ký HĐ xét duyệt thuyết minh", "Nộp BM06, chốt phiếu và lập BM07."),
    ("Thành viên HĐ nghiệm thu", "Xem hồ sơ nghiệm thu và nộp BM11."),
    ("Thư ký HĐ nghiệm thu", "Nộp BM11, chốt phiếu và lập BM12."),
]


def set_font(run, size=13, bold=None, color=TEXT):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    tr_pr.append(el)


def set_repeat_table_style(table, widths=None, header=True):
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "9FBAD0")
    for r_idx, row in enumerate(table.rows):
        no_split(row)
        if r_idx == 0 and header:
            repeat_header(row)
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cell)
            if widths and c_idx < len(widths):
                cell.width = Inches(widths[c_idx])
            if r_idx == 0 and header:
                shade(cell, BLUE)
            elif r_idx % 2 == 0:
                shade(cell, LIGHT_BLUE)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if r_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_font(run, 13, bold=(r_idx == 0), color=RGBColor(255, 255, 255) if r_idx == 0 else TEXT)


def set_cell_text(cell, text, bold=False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    set_font(r, 13, bold=bold)
    return p


def add_table(doc, headers, rows, widths=None, center_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            set_cell_text(row.cells[i], value, align=WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT)
    set_repeat_table_style(table, widths=widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Paragraph")
    p.style = doc.styles["List Paragraph"]
    p.paragraph_format.left_indent = Cm(0.63)
    p.paragraph_format.first_line_indent = Cm(-0.32)
    r = p.add_run("• ")
    set_font(r, 13)
    r = p.add_run(text)
    set_font(r, 13)
    return p


def add_numbered(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.first_line_indent = Cm(-0.7)
    r = p.add_run(f"{number}. ")
    set_font(r, 13, bold=True)
    r = p.add_run(text)
    set_font(r, 13)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    sizes = {1: 17, 2: 15, 3: 14}
    set_font(r, sizes[level], bold=True, color=RGBColor(31, 78, 121))
    return p


def add_para(doc, text="", bold_prefix=None, italic=False, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        a, b = text[: len(bold_prefix)], text[len(bold_prefix):]
        r = p.add_run(a)
        set_font(r, 13, bold=True)
        r = p.add_run(b)
        set_font(r, 13)
    else:
        r = p.add_run(text)
        set_font(r, 13)
        r.italic = italic
    return p


def add_toc_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "Mục lục sẽ được cập nhật khi mở tài liệu trong Microsoft Word."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, txt, fld_end])
    set_font(run, 13)


def set_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)


def set_portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = Cm(3.175)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.header_distance = Cm(1.27)
    section.footer_distance = Cm(1.27)


def load_sources():
    src = Document(UC_DOCX)
    overview = []
    for row in src.tables[0].rows[1:]:
        vals = [c.text.strip() for c in row.cells]
        if vals[0].startswith("UCTQ-"):
            overview.append({"code": vals[0], "name": vals[1], "actor": vals[2], "explain": vals[3]})

    modules = {}
    uc_to_module = {}
    in_scope_codes = []
    for mid, name, table_idx in MODULES:
        items = []
        for row in src.tables[table_idx].rows[1:]:
            vals = [c.text.strip() for c in row.cells]
            if vals[0].startswith("UC-"):
                item = {"code": vals[0], "name": vals[1], "actor": vals[2], "explain": vals[3], "result": vals[4], "uctq": vals[5]}
                items.append(item)
                in_scope_codes.append(vals[0])
                uc_to_module[vals[0]] = mid
        modules[mid] = items

    pending_vals = [c.text.strip() for c in src.tables[15].rows[1].cells]
    pending = {"code": pending_vals[0], "name": pending_vals[1], "actor": pending_vals[2], "explain": pending_vals[3], "result": pending_vals[4], "uctq": pending_vals[5]}

    actor_matrix = {}
    matrix_headers = [c.text.strip() for c in src.tables[16].rows[0].cells]
    for row in src.tables[16].rows[1:]:
        vals = [c.text.strip() for c in row.cells]
        actor_matrix[vals[0]] = vals[2:]

    bm_map = {}
    for line in BASELINE.read_text(encoding="utf-8").splitlines():
        if line.startswith("| UC-"):
            cells = [x.strip() for x in line.strip().strip("|").split("|")]
            if len(cells) >= 6 and re.fullmatch(r"UC-[A-Z]+-\d+", cells[0]):
                bm_map[cells[0]] = cells[-1]
    common_bm = "BM01, BM02, BM03, BM06, BM07, BM08, BM11, BM12, BM13"
    for code in ["UC-FORM-01", "UC-FORM-02", "UC-FORM-03", "UC-FORM-04", "UC-FORM-05", "UC-FORM-07"]:
        bm_map[code] = common_bm
    bm_map["UC-FORM-08"] = "BM03, BM07, BM12"
    bm_map["UC-FORM-09"] = "BM03, BM07, BM12"
    return overview, modules, uc_to_module, in_scope_codes, pending, actor_matrix, matrix_headers, bm_map


def acceptance(item):
    name = item["name"].lower()
    result = item["result"].rstrip(".")
    actor = item["actor"]
    if any(x in name for x in ["xem", "theo dõi", "tra cứu"]):
        return f"Khi {actor} truy cập đúng đối tượng và có quyền, hệ thống hiển thị {result.lower()}; dữ liệu ngoài phạm vi quyền không được cung cấp."
    if any(x in name for x in ["nộp", "gửi", "hoàn tất", "xác nhận chốt"]):
        return f"Khi {actor} cung cấp đủ dữ liệu/tệp bắt buộc và xác nhận thao tác trong thời hạn, hệ thống {result.lower()}; thao tác sai quyền, sai trạng thái hoặc quá hạn bị từ chối."
    if "trả" in name:
        return f"Khi {actor} trả đối tượng và nhập lý do bắt buộc, hệ thống {result.lower()} và thông báo cho người phải xử lý tiếp."
    if any(x in name for x in ["tạo", "lập", "cấu hình", "quản lý", "cập nhật"]):
        return f"Khi {actor} có quyền và dữ liệu hợp lệ thực hiện chức năng, hệ thống {result.lower()} và lưu dấu vết thay đổi."
    return f"Khi {actor} thực hiện chức năng trong đúng quyền và trạng thái, hệ thống {result.lower()} và ghi nhận kết quả nghiệp vụ."


def build():
    overview, module_items, uc_to_module, codes, pending, actor_matrix, matrix_headers, bm_map = load_sources()
    assert len(overview) == 19, len(overview)
    assert len(codes) == 73, len(codes)
    assert len(codes) == len(set(codes))

    doc = Document(TEMPLATE)
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)

    set_portrait(doc.sections[0])
    for style_name in ["Normal", "List Paragraph", "Caption"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = FONT
            style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
            style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
            style.font.size = Pt(13)
    for name, size in [("Heading 1", 17), ("Heading 2", 15), ("Heading 3", 14)]:
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True

    doc.core_properties.title = "Đặc tả yêu cầu chức năng hệ thống quản lý hoạt động NCKH cấp trường"
    doc.core_properties.subject = "Functional Requirements Specification"
    doc.core_properties.comments = "Bản dự thảo 0.1"

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ ĐỒNG NAI")
    set_font(r, 15, bold=True, color=RGBColor(31, 78, 121))
    doc.add_paragraph()
    for text, size in [
        ("ĐẶC TẢ YÊU CẦU CHỨC NĂNG", 22),
        ("(FUNCTIONAL REQUIREMENTS SPECIFICATION)", 15),
        ("HỆ THỐNG QUẢN LÝ HOẠT ĐỘNG", 18),
        ("NGHIÊN CỨU KHOA HỌC CẤP TRƯỜNG", 18),
        ("TRÊN NỀN TẢNG WEBSITE", 17),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(text)
        set_font(r, size, bold=True, color=RGBColor(31, 78, 121) if size >= 17 else TEXT)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Phiên bản 0.1 - Bản dự thảo")
    set_font(r, 14, bold=True)
    doc.add_paragraph()
    meta = [
        ("Mã tài liệu", "FRS-NCKH-001"), ("Phiên bản", "0.1"), ("Ngày lập", "Tháng 07 năm 2026"),
        ("Trạng thái", "Bản dự thảo"), ("Chủ đầu tư", ""), ("Đơn vị lập", ""),
        ("Tài liệu liên quan", "Quy trình thực hiện đề tài NCKH cấp trường; BM01-BM15; danh mục use case NCKH"),
        ("Lưu ý", "Các nội dung ghi Cần xác minh chưa thuộc baseline triển khai."),
    ]
    add_table(doc, ["Thông tin", "Nội dung"], meta, widths=[2.0, 4.5])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Tài liệu này thuộc phạm vi nội bộ - Không phát hành ra bên ngoài khi chưa có sự đồng ý của chủ đầu tư.")
    set_font(r, 13, bold=True)

    doc.add_page_break()
    add_heading(doc, "LỊCH SỬ THAY ĐỔI TÀI LIỆU", 1)
    add_table(doc, ["Phiên bản", "Ngày", "Nội dung thay đổi", "Người thực hiện", "Ghi chú"], [["0.1", "07/2026", "Khởi tạo bản FRS dự thảo hệ thống NCKH", "", "Chờ review và xác nhận"]], widths=[1.0, 1.0, 3.0, 1.4, 1.5], center_cols=(0, 1))
    add_heading(doc, "MỤC LỤC", 1)
    toc_p = doc.add_paragraph()
    add_toc_field(toc_p)

    doc.add_page_break()
    add_heading(doc, "CHƯƠNG 1: TỔNG QUAN TÀI LIỆU", 1)
    add_heading(doc, "1.1. Mục đích tài liệu", 2)
    add_para(doc, "Tài liệu này mô tả yêu cầu chức năng, dữ liệu, quy tắc nghiệp vụ, trạng thái, yêu cầu phi chức năng và tiêu chí nghiệm thu của hệ thống quản lý hoạt động nghiên cứu khoa học cấp trường trên nền tảng website.")
    for t in ["Làm cơ sở cho thiết kế và phát triển phần mềm.", "Làm căn cứ kiểm thử, nghiệm thu và quản lý thay đổi yêu cầu.", "Bảo đảm truy vết từ quy trình, biểu mẫu và use case đến chức năng triển khai."]:
        add_bullet(doc, t)

    add_heading(doc, "1.2. Phạm vi hệ thống", 2)
    add_para(doc, "Hệ thống hỗ trợ quản lý quy trình từ mở đợt, đăng ký, xét duyệt hồ sơ, xét duyệt thuyết minh, theo dõi thực hiện, nghiệm thu đến lưu/công bố tài liệu hoàn tất. Hệ thống quản lý dữ liệu, tệp, trạng thái, phân quyền, thông báo và dấu vết xử lý; các hoạt động họp, ký quyết định, ký hợp đồng và tài chính diễn ra bên ngoài hệ thống.")
    for _, name, _ in MODULES:
        add_bullet(doc, name)
    add_para(doc, "Ngoài phạm vi:", bold_prefix="Ngoài phạm vi:")
    for t in ["Tạo lịch, bắt đầu, hoãn hoặc tổ chức cuộc họp Hội đồng trên hệ thống.", "Ký quyết định, ký hợp đồng, lập/xử lý/ký BM14 và xử lý chữ ký Hiệu trưởng.", "Tạm ứng, thanh toán, quyết toán, hoàn trả hoặc phê duyệt chứng từ kế toán.", "Tích hợp email, SMS, thư viện, tài chính-kế toán hoặc dịch vụ ký số bên ngoài trong phiên bản đầu.", "Triển khai ứng dụng kết quả và nghiệp vụ lưu hồ sơ chuyên môn sau bước công nhận."]:
        add_bullet(doc, t)

    add_heading(doc, "1.3. Yêu cầu nghiệp vụ cốt lõi", 2)
    core = [
        ("BR-CORE-01", "Hồ sơ và biểu mẫu chỉ có hiệu lực theo đúng phiên bản, chữ ký/tệp và mốc Nộp đã được quy định."),
        ("BR-CORE-02", "Quyền truy cập được xác định theo actor, vai trò, Hội đồng/giai đoạn, đối tượng dữ liệu và trạng thái."),
        ("BR-CORE-03", "Trạng thái và các thao tác quan trọng phải được lưu vết; audit log không được sửa/xóa bởi chức năng nghiệp vụ."),
        ("BR-CORE-04", "Tài liệu lập/ký ngoài hệ thống chỉ được lưu và công bố; hệ thống không tự giả lập quy trình ký ngoài phạm vi."),
    ]
    add_table(doc, ["Mã", "Yêu cầu", "Nơi đặc tả"], [[a, b, "Chương 2-4"] for a, b in core], widths=[1.2, 4.2, 1.4], center_cols=(0, 2))

    add_heading(doc, "1.4. Đối tượng sử dụng", 2)
    add_table(doc, ["Nhóm người dùng", "Vai trò", "Phân quyền chính"], [[a, a, b] for a, b in ACTORS], widths=[1.9, 1.8, 3.3])
    add_para(doc, "Chủ nhiệm đề tài là vai trò nghiệp vụ bao quát: trong từng hồ sơ cụ thể, vai trò này do Giảng viên hoặc Sinh viên đảm nhiệm. Thành viên nhóm nghiên cứu ngoài Chủ nhiệm chỉ là dữ liệu hồ sơ, không phải actor đăng nhập.")

    add_heading(doc, "1.5. Định nghĩa và từ viết tắt", 2)
    terms = [
        ("NCKH", "Nghiên cứu khoa học"), ("P.KHCN", "Phòng/Cán bộ phụ trách Khoa học và Công nghệ"),
        ("UCTQ", "Use case tổng quan"), ("UCCT", "Use case chi tiết/yêu cầu chức năng"),
        ("BM01-BM15", "Bộ biểu mẫu thuộc quy trình thực hiện đề tài NCKH cấp trường"),
        ("Hội đồng", "Hội đồng riêng được thành lập cho từng giai đoạn xét duyệt hoặc nghiệm thu"),
        ("Actor", "Vai trò nghiệp vụ trực tiếp tương tác với hệ thống"),
        ("Audit log", "Nhật ký chỉ thêm mới dùng để truy vết thao tác quan trọng"),
        ("Bản chính thức", "Phiên bản dữ liệu/tệp đã được người có quyền bấm Nộp hoặc xác nhận"),
        ("Cần xác minh", "Thông tin chưa đủ căn cứ để đưa vào baseline triển khai"),
    ]
    add_table(doc, ["Thuật ngữ/Viết tắt", "Định nghĩa"], terms, widths=[2.0, 4.8])

    add_heading(doc, "1.6. Tài liệu tham chiếu", 2)
    for t in ["FRS_Mau.docx - mẫu cấu trúc và trình bày.", "P.KHCN_Quy trình thực hiện đề tài NCKH cấp trường và lưu đồ PPTX.", "Bộ biểu mẫu BM01-BM15 của quy trình.", "Danh mục 19 use case tổng quan, 74 use case chi tiết và ma trận actor.", "Các memlog brainstorming đã chốt nghiệp vụ đến ngày 19/07/2026."]:
        add_bullet(doc, t)

    doc.add_page_break()
    add_heading(doc, "CHƯƠNG 2: USER STORY, LUỒNG XỬ LÝ VÀ TRẠNG THÁI", 1)
    add_heading(doc, "2.1. Bộ user story", 2)
    us_rows = []
    uctq_modules = defaultdict(set)
    for mid, items in module_items.items():
        for item in items:
            for code in re.findall(r"UCTQ-\d+", item["uctq"]):
                uctq_modules[code].add(mid)
    for i, u in enumerate(overview, start=1):
        story = f"Là {u['actor']}, tôi muốn {u['name'].lower()} để {u['explain'].rstrip('.').lower()}."
        us_rows.append((f"US-{i:02d}", u["code"], u["actor"], story, ", ".join(sorted(uctq_modules[u["code"]])) or "Chờ xác minh"))
    add_table(doc, ["Mã US", "UCTQ", "Tác nhân", "User story", "Module"], us_rows, widths=[0.8, 0.9, 1.5, 3.1, 1.0], center_cols=(0, 1, 4))

    add_heading(doc, "2.2. Luồng xử lý tổng thể", 2)
    flow = [
        ("1", "Đăng ký đề tài", "Giảng viên/Sinh viên; Giảng viên hướng dẫn/Trưởng Khoa", "Tạo, ký, nộp BM01; duyệt hoặc trả sửa tuyến đầu", "BM01", "Hồ sơ đủ điều kiện xét duyệt sơ bộ"),
        ("2", "Xét duyệt đề xuất sơ bộ", "P.KHCN; Hội đồng; Thư ký", "Tạo Hội đồng, nộp BM02, lập/xác nhận BM03", "BM02-BM03", "Đạt hoặc không đạt xét duyệt hồ sơ"),
        ("3", "Nộp thuyết minh", "Chủ nhiệm đề tài", "Tải tệp thuyết minh hoàn chỉnh", "BM04", "Thuyết minh sẵn sàng xét duyệt"),
        ("4", "Phê duyệt thuyết minh", "P.KHCN; Hội đồng; Thư ký", "Đăng BM05, nộp BM06, lập/xác nhận BM07, lưu hợp đồng", "BM05-BM07", "Đang thực hiện hoặc Không thực hiện"),
        ("5", "Thực hiện/báo cáo tiến độ", "Chủ nhiệm đề tài", "Lập và nộp báo cáo giữa kỳ", "BM08", "Báo cáo tiến độ được lưu"),
        ("6", "Nộp hồ sơ và nghiệm thu", "Chủ nhiệm; P.KHCN; Hội đồng; Thư ký", "Nộp BM09, đăng BM10, nộp BM11, lập/xác nhận BM12", "BM09-BM12", "Kết quả nghiệm thu"),
        ("7", "Chỉnh sửa sau nghiệm thu", "Chủ nhiệm; P.KHCN", "Nộp BM13 và lưu BM14 hoàn chỉnh", "BM13-BM14", "Hồ sơ hoàn tất sau nghiệm thu"),
        ("8", "Công nhận kết quả", "P.KHCN", "Đăng và lưu quyết định công nhận", "BM15", "Kết quả được công bố theo quyền"),
        ("9", "Ứng dụng và lưu hồ sơ", "Đơn vị liên quan", "Thực hiện bên ngoài phạm vi xử lý chi tiết", "-", "Ngoài phạm vi phiên bản đầu"),
    ]
    add_table(doc, ["Bước", "Giai đoạn", "Actor", "Hành động chính", "BM", "Kết quả"], flow, widths=[0.5, 1.3, 1.5, 2.2, 0.8, 1.5], center_cols=(0, 4))

    add_heading(doc, "2.3. Mô hình trạng thái", 2)
    state_sets = [
        ("2.3.1. Trạng thái đợt đăng ký", [
            ("DOT-NHAP", "Nháp", "P.KHCN đang cấu hình; người đăng ký chưa nhìn thấy.", "P.KHCN tạo đợt."),
            ("DOT-CONG-BO", "Đã công bố - chưa mở", "Người dùng thấy thông tin nhưng chưa được nộp.", "P.KHCN công bố trước ngày bắt đầu."),
            ("DOT-DANG-MO", "Đang mở", "Hệ thống đang nhận hồ sơ theo quyền và điều kiện.", "Hệ thống tự chuyển khi đến thời gian mở."),
            ("DOT-DA-DONG", "Đã đóng", "Không nhận hồ sơ mới hoặc nộp lại.", "Hệ thống tự chuyển khi hết hạn."),
        ]),
        ("2.3.2. Trạng thái hồ sơ đăng ký", [
            ("HS-NHAP", "Nháp", "Người đăng ký đang hoàn thiện hồ sơ/BM01.", "Tạo hồ sơ."),
            ("HS-CHO-DUYET", "Chờ duyệt cấp đầu", "Đã nộp, chờ Giảng viên hướng dẫn hoặc Trưởng Khoa.", "Nộp BM01 đã ký."),
            ("HS-TRA-SUA", "Trả chỉnh sửa", "Hồ sơ quay lại kèm lý do bắt buộc.", "Actor cấp đầu trả sửa."),
            ("HS-CHO-HD", "Chờ Hội đồng xét duyệt hồ sơ", "Đã qua tuyến đầu và thuộc tập đủ điều kiện.", "Actor cấp đầu duyệt."),
            ("HS-DAT", "Đạt xét duyệt hồ sơ", "Được chuyển sang giai đoạn nộp thuyết minh.", "Hội đồng kết luận đạt và biên bản được xác nhận."),
            ("HS-KHONG-DAT", "Không đạt xét duyệt hồ sơ", "Hồ sơ kết thúc, không tạo đề tài tiếp tục.", "Hội đồng kết luận không đạt."),
            ("HS-QUA-HAN", "Quá hạn", "Không được tiếp tục nộp/nộp lại hoặc duyệt tuyến đầu.", "Đợt đóng khi hồ sơ còn chờ xử lý."),
            ("HS-KHONG-CHON", "Không được chọn", "Đề tài giao trực tiếp đã thuộc hồ sơ được duyệt sớm hơn.", "Trưởng Khoa duyệt một hồ sơ cùng đề tài."),
        ]),
        ("2.3.3. Trạng thái đề tài NCKH", [
            ("DT-CHO-TM", "Chờ nộp thuyết minh", "Hồ sơ đã đạt và chờ BM04.", "Hồ sơ đạt xét duyệt sơ bộ."),
            ("DT-CHO-XD-TM", "Chờ xét duyệt thuyết minh", "BM04 đã nộp và chờ Hội đồng.", "Chủ nhiệm tải BM04."),
            ("DT-DANG-TH", "Đang thực hiện", "Được phép triển khai nghiên cứu.", "P.KHCN xác nhận BM07 kết luận thực hiện."),
            ("DT-CHO-NT", "Chờ nghiệm thu", "Đã nộp hồ sơ tổng kết để đánh giá.", "Chủ nhiệm nộp hồ sơ nghiệm thu."),
            ("DT-DA-NT", "Đã nghiệm thu", "Đề tài đã có kết quả nghiệm thu được xác nhận.", "P.KHCN xác nhận BM12 đạt."),
            ("DT-KHONG-TH", "Không thực hiện", "Dừng sau xét duyệt thuyết minh.", "BM07 kết luận không thực hiện."),
            ("DT-KHONG-DAT-NT", "Không đạt nghiệm thu", "Đề tài kết thúc với kết quả không đạt.", "BM12 kết luận không đạt."),
        ]),
        ("2.3.4. Trạng thái Hội đồng", [
            ("HD-THANH-LAP", "Đã thành lập", "Hội đồng và danh sách thành viên đã được tạo.", "P.KHCN tạo Hội đồng."),
            ("HD-DANH-GIA", "Đang đánh giá", "Thành viên xem tài liệu và nộp phiếu.", "Hội đồng được giao nhiệm vụ/hạn."),
            ("HD-LAP-BB", "Đang họp/lập biên bản", "Thư ký chốt phiếu và lập biên bản.", "Thư ký xác nhận bắt đầu lập biên bản."),
            ("HD-HOAN-TAT", "Đã hoàn tất", "Biên bản/kết quả đã được P.KHCN xác nhận.", "P.KHCN xác nhận biên bản."),
            ("HD-GIAI-TAN", "Đã giải tán", "Chỉ còn quyền tra cứu lịch sử; không tải tài liệu.", "Đợt kết thúc."),
        ]),
    ]
    for title, rows in state_sets:
        add_heading(doc, title, 3)
        add_table(doc, ["Mã trạng thái", "Tên trạng thái", "Ý nghĩa", "Điều kiện chuyển"], rows, widths=[1.3, 1.5, 2.8, 2.2], center_cols=(0, 1))

    add_heading(doc, "2.3.5. Trạng thái biểu mẫu", 3)
    form_states = [
        ("BM do Chủ nhiệm nộp", "BM01/BM04/BM08/BM09/BM13", "Nháp → Đã nộp → Được chấp nhận / Trả chỉnh sửa / Quá hạn"),
        ("Phiếu Hội đồng", "BM02/BM06/BM11", "Nháp → Đã nộp / Quá hạn; đã nộp không được sửa hoặc nộp lại"),
        ("Biên bản Hội đồng", "BM03/BM07/BM12", "Nháp → Chờ P.KHCN xác nhận → Trả chỉnh sửa → Đã xác nhận"),
    ]
    add_table(doc, ["Nhóm", "Biểu mẫu", "Chuỗi trạng thái"], form_states, widths=[1.7, 1.5, 4.6])

    # Chapter 3 in landscape
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_landscape(sec)
    add_heading(doc, "CHƯƠNG 3: YÊU CẦU CHỨC NĂNG CHI TIẾT", 1)
    add_para(doc, "Chương này sử dụng trực tiếp mã UC chi tiết làm mã yêu cầu chức năng. Mọi yêu cầu trong phạm vi phiên bản đầu có ưu tiên Cao và trạng thái Đề xuất cho đến khi tài liệu được ký xác nhận.")
    for idx, (mid, mname, _) in enumerate(MODULES, start=1):
        add_heading(doc, f"3.{idx}. Module {mid} - {mname}", 2)
        items = module_items[mid]
        actors = sorted({a.strip() for x in items for a in re.split(r",| hoặc ", x["actor"]) if a.strip()})
        add_heading(doc, f"3.{idx}.1. Mô tả chức năng", 3)
        add_para(doc, f"Module {mid} nhóm {len(items)} yêu cầu liên quan đến {mname.lower()}. Actor tham gia: {'; '.join(actors)}.")
        add_heading(doc, f"3.{idx}.2. Yêu cầu chức năng", 3)
        req_rows = [[x["code"], x["name"], x["actor"], f"{x['explain']} Kết quả: {x['result']}", "Cao", "Đề xuất"] for x in items]
        add_table(doc, ["Mã YC", "Tên chức năng", "Actor", "Mô tả và kết quả", "Ưu tiên", "Trạng thái"], req_rows, widths=[1.0, 1.55, 1.55, 3.75, 0.7, 0.85], center_cols=(0, 4, 5))
        add_heading(doc, f"3.{idx}.3. Trường dữ liệu chính", 3)
        add_table(doc, ["STT", "Nhóm dữ liệu", "Mô tả"], [[str(i), f, "Dữ liệu được quản lý theo đúng quyền, trạng thái và biểu mẫu liên quan; chi tiết chưa có trong nguồn được đánh dấu Cần xác minh."] for i, f in enumerate(MODULE_FIELDS[mid], start=1)], widths=[0.5, 2.6, 6.3], center_cols=(0,))
        add_heading(doc, f"3.{idx}.4. Quy tắc nghiệp vụ", 3)
        add_table(doc, ["Mã BR", "Quy tắc", "UC liên quan"], [[f"BR-{mid}-{i:02d}", rule, ", ".join(x["code"] for x in items)] for i, rule in enumerate(MODULE_RULES[mid], start=1)], widths=[1.2, 5.8, 2.4], center_cols=(0,))
        add_heading(doc, f"3.{idx}.5. Kết quả đầu ra", 3)
        for x in items:
            add_bullet(doc, f"{x['code']}: {x['result']}")
        add_heading(doc, f"3.{idx}.6. Tiêu chí nghiệm thu", 3)
        add_table(doc, ["Mã AC", "UC", "Tiêu chí nghiệm thu"], [[f"AC-{x['code']}", x["code"], acceptance(x)] for x in items], widths=[1.4, 1.2, 6.8], center_cols=(0, 1))

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_portrait(sec)
    add_heading(doc, "CHƯƠNG 4: YÊU CẦU PHI CHỨC NĂNG", 1)
    nfr_rows = [
        ("NFR-SEC-01", "Bảo mật và phân quyền", "Người dùng chỉ được xem, tải và thao tác dữ liệu thuộc vai trò, Hội đồng/giai đoạn và đối tượng được phân công.", "Cao", "Đề xuất"),
        ("NFR-AUD-01", "Lưu vết", "Audit log được lưu trong cơ sở dữ liệu theo cơ chế chỉ thêm mới; tối thiểu gồm actor, thời điểm, hành động, đối tượng, trạng thái trước/sau và lý do.", "Cao", "Đề xuất"),
        ("NFR-DOC-01", "Tài liệu", "Hệ thống kiểm soát định dạng, phiên bản và quyền xem/tải PDF hoặc tài liệu hoàn chỉnh theo vai trò và trạng thái.", "Cao", "Đề xuất"),
        ("NFR-REL-01", "Sao lưu/khôi phục", "Dữ liệu và tệp phải được sao lưu và có khả năng khôi phục khi xảy ra sự cố; RPO/RTO cần xác minh.", "Cao", "Đề xuất"),
        ("NFR-UX-01", "Khả dụng giao diện", "Giao diện sử dụng được trên máy tính và điện thoại, hỗ trợ các trình duyệt hiện hành; danh sách phiên bản trình duyệt cần xác minh.", "Cao", "Đề xuất"),
        ("NFR-DAT-01", "Toàn vẹn dữ liệu", "Khi nhiều người đồng thời tạo/nộp hồ sơ, hệ thống không làm mất, ghi trùng hoặc gán sai dữ liệu; tải đồng thời cụ thể cần xác minh.", "Cao", "Đề xuất"),
    ]
    add_table(doc, ["Mã NFR", "Nhóm", "Yêu cầu", "Ưu tiên", "Trạng thái"], nfr_rows, widths=[1.2, 1.4, 3.8, 0.8, 1.0], center_cols=(0, 3, 4))
    add_heading(doc, "4.1. Các chỉ số cần xác minh", 2)
    for t in ["Số người dùng đồng thời và thời gian phản hồi mục tiêu.", "Dung lượng tệp tối đa và định dạng cho từng BM tải hoàn chỉnh.", "Chu kỳ sao lưu, thời gian lưu bản sao, RPO và RTO.", "Danh sách trình duyệt/phiên bản tối thiểu được hỗ trợ."]:
        add_bullet(doc, t)

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_landscape(sec)
    add_heading(doc, "CHƯƠNG 5: BẢNG TỔNG HỢP VÀ TRUY VẾT", 1)
    add_heading(doc, "5.1. Tổng hợp yêu cầu theo module", 2)
    summary = [[mid, name, len(module_items[mid]), ", ".join(x["code"] for x in module_items[mid]), "Cao", "Đề xuất"] for mid, name, _ in MODULES]
    add_table(doc, ["Module", "Tên module", "Số YC", "Mã YC", "Ưu tiên", "Trạng thái"], summary, widths=[0.7, 2.0, 0.7, 4.2, 0.8, 1.0], center_cols=(0, 2, 4, 5))
    add_para(doc, f"Tổng cộng: 14 module, {len(codes)} yêu cầu chức năng trong phạm vi và {len(nfr_rows)} yêu cầu phi chức năng.", bold_prefix="Tổng cộng:")

    add_heading(doc, "5.2. Ma trận UCTQ - UCCT - module - biểu mẫu", 2)
    trace_rows = []
    for mid, _, _ in MODULES:
        for x in module_items[mid]:
            trace_rows.append([x["uctq"], x["code"], x["name"], mid, bm_map.get(x["code"], "-")])
    add_table(doc, ["UCTQ", "UCCT", "Tên yêu cầu", "Module", "BM"], trace_rows, widths=[1.25, 1.1, 4.2, 0.8, 2.0], center_cols=(0, 1, 3, 4))

    add_heading(doc, "5.3. Ma trận use case chi tiết - actor", 2)
    matrix_abbr = ["GV", "SV", "TK", "P.KHCN", "TV-HS", "TK-HS", "TV-TM", "TK-TM", "TV-NT", "TK-NT"]
    matrix_rows = []
    for mid, _, _ in MODULES:
        for x in module_items[mid]:
            vals = actor_matrix.get(x["code"], [""] * 10)
            vals = vals[:10] + [""] * max(0, 10 - len(vals))
            matrix_rows.append([x["code"], x["name"]] + [("✓" if v.strip() else "") for v in vals])
    add_table(doc, ["Mã UCCT", "Use case chi tiết"] + matrix_abbr, matrix_rows, widths=[1.0, 3.2] + [0.52] * 10, center_cols=tuple(range(2, 12)))
    add_para(doc, "Chú giải: GV - Giảng viên; SV - Sinh viên; TK - Trưởng Khoa/Trưởng đơn vị; TV/TK-HS - Thành viên/Thư ký Hội đồng xét duyệt hồ sơ; TV/TK-TM - Hội đồng xét duyệt thuyết minh; TV/TK-NT - Hội đồng nghiệm thu.")

    add_heading(doc, "5.4. Ma trận BM01-BM15", 2)
    bm_rows = [
        ("BM01", "Nhập form + tải PDF đã ký", "Giảng viên/Sinh viên", "Giảng viên hướng dẫn hoặc Trưởng Khoa", "Hồ sơ đã nộp/được xử lý"),
        ("BM02", "Nhập form + tải PDF đã ký", "Thành viên/Thư ký/Chủ tịch HĐ hồ sơ", "Hệ thống khóa sau nộp", "Đã nộp hoặc Quá hạn"),
        ("BM03", "Nhập form, ký ngoài và tải lại", "Thư ký HĐ hồ sơ", "P.KHCN/Chủ tịch", "Đã xác nhận hoặc Trả chỉnh sửa"),
        ("BM04", "Tải tệp hoàn chỉnh", "Chủ nhiệm đề tài", "Hội đồng thuyết minh được xem/tải", "Đã nộp"),
        ("BM05", "Tải quyết định hoàn chỉnh", "P.KHCN", "Công bố theo quyền", "Đã công bố"),
        ("BM06", "Nhập form + tải PDF đã ký", "Thành viên/Thư ký/Chủ tịch HĐ thuyết minh", "Hệ thống khóa sau nộp", "Đã nộp hoặc Quá hạn"),
        ("BM07", "Nhập form, ảnh chữ ký hoặc PDF đã ký", "Thư ký/Chủ tịch HĐ thuyết minh", "P.KHCN rà soát/xác nhận", "Đã xác nhận hoặc Trả chỉnh sửa"),
        ("BM08", "Nhập form + tải PDF đã ký", "Chủ nhiệm đề tài", "Actor có quyền", "Đã nộp/Được chấp nhận/Trả sửa/Quá hạn"),
        ("BM09", "Tải báo cáo hoàn chỉnh", "Chủ nhiệm đề tài", "Actor có quyền", "Đã nộp"),
        ("BM10", "Tải quyết định hoàn chỉnh", "P.KHCN", "Công bố theo quyền", "Đã công bố"),
        ("BM11", "Nhập form + tải PDF đã ký", "Thành viên/Thư ký/Chủ tịch HĐ nghiệm thu", "Hệ thống khóa sau nộp", "Đã nộp hoặc Quá hạn"),
        ("BM12", "Nhập form, ký ngoài và tải lại", "Thư ký HĐ nghiệm thu", "P.KHCN/Chủ tịch", "Đã xác nhận hoặc Trả chỉnh sửa"),
        ("BM13", "Nhập form + tải PDF đã ký", "Chủ nhiệm đề tài", "Actor có quyền", "Đã nộp/Được chấp nhận/Trả sửa/Quá hạn"),
        ("BM14", "Tải bản hoàn chỉnh", "P.KHCN", "Lưu và cấp quyền xem", "Đã lưu"),
        ("BM15", "Tải quyết định hoàn chỉnh", "P.KHCN", "Công bố theo quyền", "Đã công bố"),
    ]
    add_table(doc, ["BM", "Phương thức", "Actor lập/nộp", "Actor/hệ thống xử lý", "Trạng thái kết thúc"], bm_rows, widths=[0.6, 2.0, 2.3, 2.3, 2.2], center_cols=(0,))

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_portrait(sec)
    add_heading(doc, "CHƯƠNG 6: PHỤ LỤC VÀ XÁC NHẬN", 1)
    add_heading(doc, "6.1. Ký hiệu mức độ ưu tiên", 2)
    add_table(doc, ["Ký hiệu", "Định nghĩa"], [("Cao", "Yêu cầu bắt buộc trong phạm vi phiên bản đầu."), ("Trung bình", "Yêu cầu quan trọng có thể triển khai sau."), ("Thấp", "Yêu cầu mong muốn, có thể hoãn khi hạn chế nguồn lực.")], widths=[1.5, 5.3])
    add_heading(doc, "6.2. Trạng thái yêu cầu", 2)
    add_table(doc, ["Trạng thái", "Mô tả"], [("Đề xuất", "Đã có trong dự thảo nhưng chưa được chủ đầu tư ký xác nhận."), ("Đã xác nhận", "Đã được chủ đầu tư xác nhận chính thức."), ("Đang phát triển", "Đội phát triển đang triển khai."), ("Hoàn thành", "Đã triển khai và kiểm thử thành công."), ("Cần xác minh", "Chưa đủ dữ liệu để trở thành baseline triển khai.")], widths=[1.8, 5.0])
    add_heading(doc, "6.3. Điểm giao tiếp và tích hợp", 2)
    integrations = [
        ("Tài khoản trong Trường", "Dùng tài khoản do Nhà trường quản lý; cơ chế SSO cụ thể cần xác minh."),
        ("Tài khoản ngoài Trường", "P.KHCN cấp và phân quyền cho thành viên/Thư ký Hội đồng ngoài Trường."),
        ("Email/SMS", "Ngoài phạm vi phiên bản đầu; chỉ dùng thông báo trong ứng dụng."),
        ("Thư viện", "Ngoài phạm vi phiên bản đầu."),
        ("Tài chính-kế toán", "Ngoài phạm vi phiên bản đầu."),
        ("Ký số bên ngoài", "Chưa tích hợp bắt buộc; BM07 vẫn có phương thức ảnh chữ ký hoặc PDF đã ký."),
    ]
    add_table(doc, ["Điểm giao tiếp", "Mô tả"], integrations, widths=[2.0, 4.8])
    add_heading(doc, "6.4. Danh sách vấn đề mở", 2)
    open_issues = [
        ("OPEN-01", pending["code"], pending["name"], "Chưa xác định đối tượng tra cứu, nguồn dữ liệu, điều kiện và kết quả; không tính vào 73 yêu cầu triển khai."),
        ("OPEN-02", "NFR", "Chỉ số vận hành", "Cần xác minh tải đồng thời, thời gian phản hồi, dung lượng tệp, RPO và RTO."),
        ("OPEN-03", "FORM/PDF", "Nguồn dữ liệu chính thức", "Cần xác minh cách xử lý khi dữ liệu form và PDF đã ký khác nhau."),
        ("OPEN-04", "Tài liệu", "Định dạng và dung lượng", "Cần xác minh định dạng/dung lượng tối đa theo từng BM tải hoàn chỉnh."),
        ("OPEN-05", "Phê duyệt", "Thông tin ký xác nhận FRS", "Chưa có họ tên đại diện Chủ đầu tư, Trưởng nhóm BA và người phê duyệt kỹ thuật."),
    ]
    add_table(doc, ["Mã", "Liên quan", "Nội dung", "Mô tả"], open_issues, widths=[1.0, 1.2, 1.8, 3.8], center_cols=(0, 1))
    add_heading(doc, "6.5. Xác nhận tài liệu", 2)
    add_para(doc, "Việc ký xác nhận chuyển trạng thái các yêu cầu từ Đề xuất sang Đã xác nhận. Các vấn đề mở không mặc nhiên trở thành yêu cầu triển khai cho đến khi có quyết định bổ sung.")
    add_table(doc, ["Vai trò", "Họ tên", "Ngày ký"], [("Đại diện Chủ đầu tư", "", ""), ("Trưởng nhóm BA", "", ""), ("Phê duyệt kỹ thuật", "", "")], widths=[2.3, 2.7, 1.8])

    # Normalize body/table font and field-update setting.
    for p in doc.paragraphs:
        if not p.style.name.startswith("Heading"):
            for run in p.runs:
                if run.font.size is None or run.font.size.pt <= 13.1:
                    set_font(run, 13, bold=run.bold)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_font(run, 13, bold=run.bold, color=run.font.color.rgb or TEXT)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"OVERVIEW={len(overview)} REQUIREMENTS={len(codes)} MODULES={len(MODULES)} TABLES={len(doc.tables)} SECTIONS={len(doc.sections)}")


if __name__ == "__main__":
    build()
